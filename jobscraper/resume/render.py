"""Write a TailoredResume to DOCX (and a Markdown copy).

Two ways to control the look:
  1. Your own Word template (RESUME_TEMPLATE_PATH=templates/resume_template.docx)
     using docxtpl / Jinja tags — see templates/README.md. Run
     `python -m jobscraper make-template` to get a starter template to restyle.
  2. No template: a clean single-column, ATS-friendly layout built here.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from ..config import load_settings
from .parser import Resume
from .tailor import TailoredResume

CONTACT_ORDER = ["email", "phone", "location", "linkedin", "github", "portfolio", "website"]


def safe_filename(text: str, max_len: int = 60) -> str:
    text = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE).strip()
    return re.sub(r"[\s_-]+", "_", text)[:max_len].strip("_") or "resume"


def output_path(t: TailoredResume, job_id: str = "", suffix: str = ".docx") -> Path:
    settings = load_settings()
    folder = settings.resumes_dir / date.today().isoformat()
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{safe_filename(t.company, 30)}_{safe_filename(t.job_title, 50)}"
    if job_id:
        name += f"_{safe_filename(str(job_id), 12)}"
    return folder / f"{name}{suffix}"


def template_context(t: TailoredResume) -> dict:
    r: Resume = t.resume
    contact_items = [r.contact[k] for k in CONTACT_ORDER if r.contact.get(k)]
    contact_items += [v for k, v in r.contact.items() if k not in CONTACT_ORDER]
    sections = []
    for s in r.sections:
        sections.append({"name": s.name, "paragraph": s.paragraph, "bullets": s.bullets,
                         "entries": [e.as_dict() for e in s.entries]})
    by_name = {s["name"].lower(): s for s in sections}
    return {
        "name": r.name,
        "contact": r.contact,
        "contact_line": "  |  ".join(contact_items),
        "summary": r.summary,
        "skills": [g.as_dict() for g in r.skills],
        "key_skills": ", ".join(t.matched_keywords[:15]),
        "sections": sections,
        "experience": by_name.get("experience", {}).get("entries", []),
        "projects": by_name.get("projects", {}).get("entries", []),
        "education": by_name.get("education", {}).get("entries", []),
        "certifications": by_name.get("certifications", {}).get("bullets", []),
        "achievements": by_name.get("achievements", {}).get("bullets", []),
        "job_title": t.job_title,
        "company": t.company,
    }


# ── built-in layout ─────────────────────────────────────────────────────────

def _build_docx(t: TailoredResume, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    ctx = template_context(t)
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(36)
        s.left_margin = s.right_margin = Pt(42)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(1)

    def heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        # bottom border under section headings
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)

    def bullet(text: str):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        p.add_run(text)

    def entry(e: dict):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        r = p.add_run(e["title"])
        r.bold = True
        if e["organization"]:
            p.add_run(f"  |  {e['organization']}")
        right = "  |  ".join(x for x in (e["location"], e["dates"]) if x)
        if right:
            p.add_run(f"  |  {right}").italic = True
        for x in e["extra"]:
            doc.add_paragraph(x).paragraph_format.space_after = Pt(0)
        for b in e["bullets"]:
            bullet(b)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(ctx["name"])
    r.bold = True
    r.font.size = Pt(18)
    if ctx["contact_line"]:
        p = doc.add_paragraph(ctx["contact_line"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if ctx["summary"]:
        heading("Summary")
        doc.add_paragraph(ctx["summary"])

    if ctx["skills"] or ctx["key_skills"]:
        heading("Skills")
        if ctx["key_skills"]:
            p = doc.add_paragraph()
            p.add_run("Key skills for this role: ").bold = True
            p.add_run(ctx["key_skills"])
        for g in ctx["skills"]:
            p = doc.add_paragraph()
            p.add_run(f"{g['category']}: ").bold = True
            p.add_run(g["text"])

    for s in ctx["sections"]:
        heading(s["name"])
        if s["paragraph"]:
            doc.add_paragraph(s["paragraph"])
        for e in s["entries"]:
            entry(e)
        for b in s["bullets"]:
            bullet(b)

    doc.core_properties.title = f"{ctx['name']} - {t.job_title} - {t.company}"
    doc.core_properties.author = ctx["name"]
    doc.save(str(path))


def _render_template(t: TailoredResume, template: Path, path: Path) -> None:
    from docxtpl import DocxTemplate

    tpl = DocxTemplate(str(template))
    tpl.render(template_context(t), autoescape=True)
    tpl.save(str(path))


def to_markdown(t: TailoredResume) -> str:
    ctx = template_context(t)
    out = [f"# {ctx['name']}", ctx["contact_line"], ""]
    if ctx["summary"]:
        out += ["## Summary", ctx["summary"], ""]
    out.append("## Skills")
    if ctx["key_skills"]:
        out.append(f"**Key skills for this role:** {ctx['key_skills']}")
    out += [f"- **{g['category']}:** {g['text']}" for g in ctx["skills"]] + [""]
    for s in ctx["sections"]:
        out.append(f"## {s['name']}")
        if s["paragraph"]:
            out.append(s["paragraph"])
        for e in s["entries"]:
            bits = " | ".join(x for x in (e["organization"], e["location"], e["dates"], *e["extra"]) if x)
            out.append(f"### {e['title']}" + (f" | {bits}" if bits else ""))
            out += [f"- {b}" for b in e["bullets"]]
        out += [f"- {b}" for b in s["bullets"]] + [""]
    return "\n".join(out).strip() + "\n"


def render_resume(t: TailoredResume, job_id: str = "") -> Path:
    """Write the tailored resume; returns the .docx path (a .md copy sits next to it)."""
    path = output_path(t, job_id)
    template = load_settings().resume_template_path
    if template and template.is_file():
        _render_template(t, template, path)
    else:
        _build_docx(t, path)
    path.with_suffix(".md").write_text(to_markdown(t), encoding="utf-8")
    return path


def make_starter_template(path: Path) -> Path:
    """Create a docxtpl starter template the user can restyle in Word."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    p = doc.add_paragraph()
    p.add_run("{{ name }}").bold = True
    doc.add_paragraph("{{ contact_line }}")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph("{{ summary }}")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Key skills for this role: {{ key_skills }}")
    doc.add_paragraph("{%p for g in skills %}")
    doc.add_paragraph("{{ g.category }}: {{ g.text }}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p for s in sections %}")
    doc.add_heading("{{ s.name }}", level=2)
    doc.add_paragraph("{% if s.paragraph %}{{ s.paragraph }}{% endif %}")
    doc.add_paragraph("{%p for e in s.entries %}")
    doc.add_paragraph("{{ e.title }}{% if e.organization %} | {{ e.organization }}{% endif %}"
                      "{% if e.location %} | {{ e.location }}{% endif %}{% if e.dates %} | {{ e.dates }}{% endif %}")
    doc.add_paragraph("{%p for x in e.extra %}")
    doc.add_paragraph("{{ x }}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p for b in e.bullets %}")
    doc.add_paragraph("{{ b }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p for b in s.bullets %}")
    doc.add_paragraph("{{ b }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endfor %}")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
